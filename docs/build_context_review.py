"""Build curator_context_review.xlsx from docs/context/ and Curator.json.

Run with:  python -I build_context_review.py
"""
import json
import os
import re
import sys
from collections import Counter, defaultdict
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DOCS_DIR = os.path.dirname(os.path.abspath(__file__))
CONTEXT_DIR = os.path.join(DOCS_DIR, "context")
META_PATH = os.path.join(CONTEXT_DIR, "Curator.json")
OUT_PATH = os.path.join(DOCS_DIR, "curator_context_review.xlsx")

SKIP_FILES = {"Curator.json", ".DS_Store", ".gitkeep"}
SOURCE_EXT = {"py", "js", "css", "sql", "html"}
CONFIG_EXT = {"yaml", "yml", "toml", "json", "cfg"}
HANDOFF_DAYS = 60
SNIPPET = 150
TODAY = date.today()

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
ARCHIVE_FILL = PatternFill("solid", fgColor="F2F2F2")
CHANGEDOC_FILL = PatternFill("solid", fgColor="F0FFF0")
KEEP_FILL = PatternFill("solid", fgColor="FFFACD")
HEADER_FONT = Font(bold=True, color="FFFFFF")

COLUMNS = ["file_name", "category", "file_type", "created_at", "token_count",
           "had_duplicates", "recommended_action", "description", "keep", "notes"]
MAX_WIDTH = {"description": 60, "file_name": 30}


def collapse(text):
    return re.sub(r"\s+", " ", text).strip()


def clip(text, n=SNIPPET):
    text = collapse(text)
    return text if len(text) <= n else text[: n - 1].rstrip() + "…"


def file_ext(name):
    base = os.path.basename(name)
    return base.rsplit(".", 1)[1].lower() if "." in base.lstrip(".") else ""


def is_test(base):
    return base.endswith(".py") and (base.startswith("test_") or base.endswith("_test.py")
                                     or base == "conftest.py")


def categorize(base):
    ext = file_ext(base)
    lower = base.lower()
    if base.startswith("CHANGEDOC_"):
        return "CHANGEDOC"
    if base.startswith("DESIGN_") or re.search(r"_design(_|\.)", lower):
        return "DESIGN"
    if "handoff" in lower:
        return "handoff"
    if is_test(base):
        return "test"
    if ext in SOURCE_EXT:
        return "source_" + ext
    if ext in CONFIG_EXT or base == "Dockerfile":
        return "config"
    return "other"


def handoff_date(created, base):
    """created_at, falling back to a YYYY-MM-DD date in the filename."""
    if created:
        return created
    m = re.search(r"(\d{4}-\d{2}-\d{2})", base)
    if m:
        try:
            return datetime.strptime(m.group(1), "%Y-%m-%d").date()
        except ValueError:
            pass
    return None


def recommend(base, category, created):
    ext = file_ext(base)
    if ext in SOURCE_EXT or is_test(base):
        return "archive"
    if base.startswith("CHANGEDOC_"):
        return "keep"
    if base.startswith("DESIGN_"):
        return "review"
    if "handoff" in base.lower():
        d = handoff_date(created, base)
        return "keep" if d and (TODAY - d).days <= HANDOFF_DAYS else "review"
    return "review"


def describe_md(text):
    lines = text.splitlines()
    h1 = next((l[2:].strip() for l in lines if re.match(r"#\s", l)), "")
    para, seen_heading, in_code = [], False, False
    for line in lines:
        s = line.strip()
        if s.startswith("```"):
            in_code = not in_code
            if para:
                break
            continue
        if in_code:
            continue
        if re.match(r"#{1,6}\s", s):
            if para:
                break
            seen_heading = True
            continue
        if not seen_heading:
            continue
        if not s or re.fullmatch(r"[-*_=]{3,}", s):
            if para:
                break
            continue
        para.append(s)
    if not seen_heading:
        # No headings at all: use the first non-empty block.
        para = []
        for s in (l.strip() for l in lines):
            if s:
                para.append(s)
            elif para:
                break
    snippet = clip(re.sub(r"\*\*|__", "", " ".join(para)))
    if h1 and snippet:
        return f"{h1} — {snippet}"
    return h1 or snippet


COMMENT_PATTERNS = {
    "py": [r'^\s*[rRuU]?("""|\'\'\')(.*?)\1', r"^((?:[ \t]*#.*\n?)+)"],
    "js": [r"/\*(.*?)\*/", r"((?:^[ \t]*//.*\n?)+)"],
    "css": [r"/\*(.*?)\*/"],
    "sql": [r"((?:^[ \t]*--.*\n?)+)", r"/\*(.*?)\*/"],
    "html": [r"<!--(.*?)-->", r"\{#(.*?)#\}"],
}


def strip_markers(block):
    out = []
    for line in block.splitlines():
        s = line.strip()
        s = re.sub(r"^(#!.*|#|//|--|/\*+|\*+/|\*|<!--|-->|\{#|#\})", "", s)
        s = re.sub(r"(\*+/|-->|#\})$", "", s).strip()
        s = s.strip("=-*# ")
        if s:
            out.append(s)
    return " ".join(out)


def describe_source(text, ext):
    best = None
    for pat in COMMENT_PATTERNS[ext]:
        for m in re.finditer(pat, text, re.S | re.M):
            body = strip_markers(m.group(m.lastindex))
            if body:
                if best is None or m.start() < best[0]:
                    best = (m.start(), body)
                break
    return clip(best[1]) if best else clip(text)


def describe(base, text):
    ext = file_ext(base)
    if ext == "md":
        return describe_md(text)
    if ext in SOURCE_EXT:
        return describe_source(text, ext)
    return clip(text)


def load_meta():
    with open(META_PATH, encoding="utf-8") as f:
        entries = json.load(f)
    counts = Counter()
    latest = {}
    for e in entries:
        if not isinstance(e, dict) or not isinstance(e.get("file_name"), str):
            continue
        name = e["file_name"]
        counts[name] += 1
        if name not in latest or str(e.get("created_at", "")) > str(latest[name].get("created_at", "")):
            latest[name] = e
    return latest, counts


def list_files():
    out = []
    for root, dirs, files in os.walk(CONTEXT_DIR):
        dirs[:] = sorted(d for d in dirs if not d.startswith("."))
        for f in sorted(files):
            if f in SKIP_FILES:
                continue
            out.append(os.path.relpath(os.path.join(root, f), CONTEXT_DIR))
    return out


def parse_date(value):
    if not isinstance(value, str):
        return None
    try:
        return datetime.fromisoformat(value.replace("Z", "+00:00")[:32]).date()
    except ValueError:
        m = re.match(r"\d{4}-\d{2}-\d{2}", value)
        return datetime.strptime(m.group(0), "%Y-%m-%d").date() if m else None


GROUP_ORDER = {"CHANGEDOC": 0, "DESIGN": 1, "handoff": 2}


def group_rank(category):
    if category in GROUP_ORDER:
        return GROUP_ORDER[category]
    if category.startswith("source_") or category == "test":
        return 3
    return 4


def build_rows():
    latest, counts = load_meta()
    rows, unmatched = [], []
    for rel in list_files():
        base = os.path.basename(rel)
        meta = latest.get(rel)
        if meta is None:
            unmatched.append(rel)
        created = parse_date(meta.get("created_at")) if meta else None
        tokens = meta.get("estimated_token_count") if meta else None
        tokens = tokens if isinstance(tokens, int) else None
        with open(os.path.join(CONTEXT_DIR, rel), encoding="utf-8", errors="replace") as f:
            text = f.read()
        category = categorize(base)
        rows.append({
            "file_name": base,
            "category": category,
            "file_type": file_ext(base),
            "created_at": created,
            "token_count": tokens,
            "had_duplicates": counts.get(rel, 0) > 1,
            "recommended_action": recommend(base, category, created),
            "description": describe(base, text),
            "keep": None,
            "notes": None,
            "_rel": rel,
        })
    rows.sort(key=lambda r: (group_rank(r["category"]), r["file_name"].lower(), r["_rel"]))
    return rows, unmatched, counts


def set_text(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # never let content be parsed as a formula


def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL


def write_decisions(ws, rows):
    ws.append(COLUMNS)
    style_header(ws, len(COLUMNS))
    keep_col = COLUMNS.index("keep") + 1
    for i, r in enumerate(rows, start=2):
        fill = None
        if r["recommended_action"] == "archive":
            fill = ARCHIVE_FILL
        elif r["category"] == "CHANGEDOC":
            fill = CHANGEDOC_FILL
        for c, key in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=i, column=c)
            set_text(cell, r[key])
            if key == "created_at" and r[key]:
                cell.number_format = "yyyy-mm-dd"
            if fill:
                cell.fill = fill
        ws.cell(row=i, column=keep_col).fill = KEEP_FILL

    dv = DataValidation(type="list", formula1='"Yes,No,Review"', allow_blank=True)
    col = get_column_letter(keep_col)
    dv.add(f"{col}2:{col}{len(rows) + 1}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    for c, key in enumerate(COLUMNS, start=1):
        width = len(key)
        for r in rows:
            v = r[key]
            if v is None:
                continue
            width = max(width, 10 if isinstance(v, date) else len(str(v)))
        ws.column_dimensions[get_column_letter(c)].width = min(width + 2, MAX_WIDTH.get(key, 60))
    desc_col = COLUMNS.index("description") + 1
    for i in range(2, len(rows) + 2):
        ws.cell(row=i, column=desc_col).alignment = Alignment(wrap_text=True, vertical="top")


def write_summary(ws, rows, counts):
    bold = Font(bold=True)
    total_tokens = sum(r["token_count"] or 0 for r in rows)

    def header(row, labels):
        for c, label in enumerate(labels, start=1):
            ws.cell(row=row, column=c, value=label)
        style_header_row(row, len(labels))

    def style_header_row(row, n):
        for c in range(1, n + 1):
            ws.cell(row=row, column=c).font = HEADER_FONT
            ws.cell(row=row, column=c).fill = HEADER_FILL

    ws.cell(row=1, column=1, value="Category breakdown").font = bold
    header(2, ["category", "count", "total_tokens", "pct_of_total_tokens"])
    by_cat = defaultdict(lambda: [0, 0])
    for r in rows:
        by_cat[r["category"]][0] += 1
        by_cat[r["category"]][1] += r["token_count"] or 0
    row = 3
    for cat, (n, tok) in sorted(by_cat.items(), key=lambda kv: -kv[1][1]):
        ws.append([cat, n, tok, tok / total_tokens if total_tokens else 0])
        ws.cell(row=row, column=4).number_format = "0.0%"
        row += 1
    ws.append(["TOTAL", len(rows), total_tokens, 1 if total_tokens else 0])
    ws.cell(row=row, column=4).number_format = "0.0%"
    for c in range(1, 5):
        ws.cell(row=row, column=c).font = bold
    row += 2

    ws.cell(row=row, column=1, value="Recommendation breakdown").font = bold
    row += 1
    header(row, ["recommended_action", "count", "total_tokens", "pct_of_total_tokens"])
    row += 1
    for action in ("keep", "review", "archive"):
        subset = [r for r in rows if r["recommended_action"] == action]
        tok = sum(r["token_count"] or 0 for r in subset)
        for c, v in enumerate([action, len(subset), tok, tok / total_tokens if total_tokens else 0], 1):
            ws.cell(row=row, column=c, value=v)
        ws.cell(row=row, column=4).number_format = "0.0%"
        row += 1
    row += 1

    dup_names = [n for n, k in counts.items() if k > 1]
    ws.cell(row=row, column=1, value="Duplicates in project context (Curator.json)").font = bold
    row += 1
    header(row, ["metric", "value"])
    row += 1
    stats = [
        ("files with duplicates", len(dup_names)),
        ("total entries for those files", sum(counts[n] for n in dup_names)),
        ("redundant (extra) entries", sum(counts[n] - 1 for n in dup_names)),
        ("total entries in Curator.json", sum(counts.values())),
        ("unique filenames in Curator.json", len(counts)),
    ]
    for label, v in stats:
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=v)
        row += 1

    for col, width in zip("ABCD", (34, 10, 14, 20)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A3"


def main():
    rows, unmatched, counts = build_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "Decisions"
    write_decisions(ws, rows)
    write_summary(wb.create_sheet("Summary"), rows, counts)
    wb.save(OUT_PATH)

    actions = Counter(r["recommended_action"] for r in rows)
    print(f"Wrote {OUT_PATH}")
    print(f"Rows: {len(rows)}  (keep {actions['keep']}, review {actions['review']}, archive {actions['archive']})")
    print(f"Files with no match in Curator.json: {len(unmatched)}")
    for rel in unmatched:
        print(f"  - {rel}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
